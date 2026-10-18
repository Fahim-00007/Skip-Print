"""
Excel file management router.
Handles uploads, sheet inspection, spreadsheet grid previews, sample loading, and live cell editing.
"""
import os
import shutil
import uuid
import openpyxl
from fastapi import APIRouter, UploadFile, File, HTTPException, Query
from backend.services.excel_reader import inspect_workbook, get_sheet_preview
from backend.models import BatchCellUpdateRequest, CellUpdateItem

router = APIRouter(prefix="/api/excel", tags=["Excel"])

STORAGE_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "storage", "uploads")
OUTPUTS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "storage", "outputs")
SAMPLES_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "samples")
os.makedirs(STORAGE_DIR, exist_ok=True)
os.makedirs(OUTPUTS_DIR, exist_ok=True)

def find_file_path(file_id: str):
    if file_id == "sample-invoice":
        sample_target = os.path.join(STORAGE_DIR, "sample-invoice_sample_invoice.xlsx")
        if not os.path.exists(sample_target):
            orig_sample = os.path.join(SAMPLES_DIR, "sample_invoice.xlsx")
            if os.path.exists(orig_sample):
                shutil.copy2(orig_sample, sample_target)
        return sample_target

    for fname in os.listdir(STORAGE_DIR):
        if fname.startswith(file_id):
            return os.path.join(STORAGE_DIR, fname)
    return None

def parse_cell_value(val_str: str):
    """
    Parses cell input: converts numeric strings to numbers, preserves formulas and text.
    """
    if val_str is None:
        return ""
    val_str = str(val_str).strip()
    if val_str.startswith("="):
        return val_str
    # Check integer
    try:
        if val_str.isdigit() or (val_str.startswith("-") and val_str[1:].isdigit()):
            return int(val_str)
    except Exception:
        pass
    # Check float (e.g. 150.00 or $150.00)
    cleaned = val_str.replace("$", "").replace(",", "").strip()
    try:
        return float(cleaned)
    except Exception:
        pass
    return val_str

@router.post("/upload")
async def upload_excel(file: UploadFile = File(...)):
    filename = file.filename or "uploaded_bill.xlsx"
    ext = os.path.splitext(filename)[1].lower()

    if ext not in [".xlsx", ".xls", ".xlsm"]:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file format '{ext}'. Please upload an Excel file (.xlsx or .xls)."
        )

    file_id = str(uuid.uuid4())
    safe_name = f"{file_id}_{os.path.basename(filename)}"
    dest_path = os.path.join(STORAGE_DIR, safe_name)

    try:
        with open(dest_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to save uploaded file: {str(e)}")

    file_size = os.path.getsize(dest_path)

    try:
        overview = inspect_workbook(dest_path)
        first_sheet = overview["active_sheet"] or overview["sheet_names"][0]
        preview_grid = get_sheet_preview(dest_path, sheet_name=first_sheet)
    except Exception as e:
        raise HTTPException(
            status_code=422,
            detail=f"Error reading Excel content: {str(e)}. Please check if the file is password-protected or corrupt."
        )

    return {
        "file_id": file_id,
        "filename": filename,
        "file_path": dest_path,
        "file_size": file_size,
        "sheet_names": overview["sheet_names"],
        "active_sheet": first_sheet,
        "sheet_count": overview["sheet_count"],
        "preview_grid": preview_grid
    }

@router.get("/{file_id}/sheet-preview")
async def get_sheet_data(file_id: str, sheet: str = Query(None)):
    target_path = find_file_path(file_id)
    if not target_path or not os.path.exists(target_path):
        raise HTTPException(status_code=404, detail="Excel file not found.")

    try:
        preview = get_sheet_preview(target_path, sheet_name=sheet)
        return preview
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate sheet preview: {str(e)}")

@router.post("/{file_id}/update-cells")
async def update_sheet_cells(file_id: str, req: BatchCellUpdateRequest):
    target_path = find_file_path(file_id)
    if not target_path or not os.path.exists(target_path):
        raise HTTPException(status_code=404, detail="Excel file not found.")

    ext = os.path.splitext(target_path)[1].lower()
    if ext not in [".xlsx", ".xlsm"]:
        raise HTTPException(status_code=400, detail="In-browser cell editing is supported for .xlsx files.")

    try:
        wb = openpyxl.load_workbook(target_path)
        target_sheet = req.sheet_name if (req.sheet_name and req.sheet_name in wb.sheetnames) else wb.active.title
        ws = wb[target_sheet]

        for item in req.updates:
            parsed = parse_cell_value(item.value)
            if item.coordinate:
                ws[item.coordinate] = parsed
            elif item.row and item.col:
                ws.cell(row=item.row, column=item.col, value=parsed)

        wb.save(target_path)
        wb.close()

        # Invalidate any cached raw PDFs for this file
        for fname in os.listdir(OUTPUTS_DIR):
            if fname.startswith(f"raw_{file_id}") or fname.startswith(f"preview_{file_id}"):
                try:
                    os.remove(os.path.join(OUTPUTS_DIR, fname))
                except Exception:
                    pass

        # Return updated grid preview
        updated_grid = get_sheet_preview(target_path, sheet_name=target_sheet)
        return {
            "success": True,
            "sheet_name": target_sheet,
            "preview_grid": updated_grid
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to save edited cells: {str(e)}")

@router.post("/load-sample")
async def load_sample_invoice():
    sample_src = os.path.join(SAMPLES_DIR, "sample_invoice.xlsx")
    if not os.path.exists(sample_src):
        from backend.samples.generate_samples import create_sample_excel
        sample_src = create_sample_excel()

    file_id = "sample-invoice"
    dest_path = os.path.join(STORAGE_DIR, f"{file_id}_sample_invoice.xlsx")
    shutil.copy2(sample_src, dest_path)

    overview = inspect_workbook(dest_path)
    preview = get_sheet_preview(dest_path, sheet_name=overview["active_sheet"])

    return {
        "file_id": file_id,
        "filename": "sample_invoice.xlsx",
        "file_path": dest_path,
        "file_size": os.path.getsize(dest_path),
        "sheet_names": overview["sheet_names"],
        "active_sheet": overview["active_sheet"],
        "sheet_count": overview["sheet_count"],
        "preview_grid": preview
    }
