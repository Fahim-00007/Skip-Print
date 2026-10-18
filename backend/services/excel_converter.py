"""
Excel-to-PDF Conversion Service.
Provides primary native COM conversion (Excel 100% faithful vector export)
and an automated pure-Python ReportLab fallback converter.
"""
import os
import sys
import tempfile
import openpyxl
from reportlab.lib.pagesizes import A4, letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def convert_excel_to_pdf(excel_path: str, output_pdf_path: str, sheet_name: str = None, engine: str = "auto") -> dict:
    """
    Converts the specified worksheet of an Excel file to PDF.
    Engine options: 'auto', 'com', 'fallback'.
    """
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    os.makedirs(os.path.dirname(os.path.abspath(output_pdf_path)), exist_ok=True)

    use_com = (engine in ["auto", "com"])
    com_error = None

    if use_com and sys.platform == "win32":
        try:
            return convert_via_com(excel_path, output_pdf_path, sheet_name)
        except Exception as e:
            com_error = str(e)
            print(f"COM conversion encountered an issue: {com_error}. Falling back to pure Python engine...")

    # Fallback to pure Python converter
    return convert_via_python(excel_path, output_pdf_path, sheet_name, warning=com_error)

def convert_via_com(excel_path: str, output_pdf_path: str, sheet_name: str = None) -> dict:
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    excel = None
    wb = None
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.ScreenUpdating = False

        abs_excel = os.path.abspath(excel_path)
        abs_output = os.path.abspath(output_pdf_path)

        if os.path.exists(abs_output):
            os.remove(abs_output)

        wb = excel.Workbooks.Open(abs_excel, ReadOnly=True, UpdateLinks=False)

        target_obj = None
        if sheet_name:
            for s in wb.Sheets:
                if s.Name == sheet_name:
                    target_obj = s
                    break

        if not target_obj:
            target_obj = wb.ActiveSheet

        # Configure page setup to ensure it fits cleanly
        try:
            target_obj.PageSetup.Orientation = 1  # 1 = xlPortrait, 2 = xlLandscape
            target_obj.PageSetup.Zoom = False
            target_obj.PageSetup.FitToPagesWide = 1
            target_obj.PageSetup.FitToPagesTall = False
        except Exception:
            pass

        # 0 = xlTypePDF
        target_obj.ExportAsFixedFormat(0, abs_output, Quality=0, IncludeDocProperties=True, IgnorePrintAreas=False)

        if not os.path.exists(abs_output) or os.path.getsize(abs_output) == 0:
            raise RuntimeError("Excel COM exported an empty PDF file.")

        return {
            "success": True,
            "engine": "Excel COM (Native Vector Fidelity)",
            "output_path": abs_output,
            "warning": None
        }

    finally:
        if wb:
            try:
                wb.Close(SaveChanges=False)
            except Exception:
                pass
        if excel:
            try:
                excel.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()

def convert_via_python(excel_path: str, output_pdf_path: str, sheet_name: str = None, warning: str = None) -> dict:
    """
    Pure Python fallback converter using openpyxl and ReportLab.
    """
    ext = os.path.splitext(excel_path)[1].lower()
    abs_output = os.path.abspath(output_pdf_path)

    table_data = []
    col_widths = []
    
    if ext in [".xlsx", ".xlsm"]:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
        ws = wb[sheet_name] if (sheet_name and sheet_name in wb.sheetnames) else wb.active

        max_r = ws.max_row or 1
        max_c = ws.max_column or 1

        # Calculate rough column widths in points (72 pt / inch)
        for c in range(1, max_c + 1):
            from openpyxl.utils import get_column_letter
            col_letter = get_column_letter(c)
            dim = ws.column_dimensions.get(col_letter)
            w = (dim.width if dim and dim.width else 10) * 6.5
            col_widths.append(min(max(w, 25), 180))

        styles = getSampleStyleSheet()
        normal_style = styles["Normal"]
        normal_style.fontSize = 8
        normal_style.leading = 10

        bold_style = ParagraphStyle('BoldStyle', parent=normal_style, fontName='Helvetica-Bold')

        for r in range(1, max_r + 1):
            row_cells = []
            for c in range(1, max_c + 1):
                cell = ws.cell(row=r, column=c)
                val = cell.value
                val_str = ""
                if val is not None:
                    if isinstance(val, float):
                        val_str = f"{val:,.2f}" if (abs(val) >= 0.01 and '.' in str(val)) else str(val)
                    else:
                        val_str = str(val)

                is_bold = bool(cell.font and cell.font.bold)
                p_style = bold_style if is_bold else normal_style
                row_cells.append(Paragraph(val_str.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), p_style))
            table_data.append(row_cells)

        wb.close()
    else:
        # Minimal legacy .xls support via xlrd
        import xlrd
        book = xlrd.open_workbook(excel_path)
        sheet = book.sheet_by_name(sheet_name) if (sheet_name and sheet_name in book.sheet_names()) else book.sheet_by_index(0)
        
        styles = getSampleStyleSheet()
        normal_style = styles["Normal"]
        normal_style.fontSize = 8
        normal_style.leading = 10

        for r in range(sheet.nrows):
            row_cells = []
            for c in range(sheet.ncols):
                val = sheet.cell_value(r, c)
                val_str = str(val) if val is not None else ""
                row_cells.append(Paragraph(val_str.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"), normal_style))
            table_data.append(row_cells)
            col_widths = [60] * sheet.ncols

    # Render ReportLab document
    doc = SimpleDocTemplate(
        abs_output,
        pagesize=A4,
        leftMargin=20,
        rightMargin=20,
        topMargin=20,
        bottomMargin=20
    )

    t = Table(table_data, colWidths=col_widths if col_widths else None)
    t.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
    ]))

    doc.build([t])

    return {
        "success": True,
        "engine": "ReportLab Fallback Engine",
        "output_path": abs_output,
        "warning": f"Used fallback engine ({warning})" if warning else None
    }
