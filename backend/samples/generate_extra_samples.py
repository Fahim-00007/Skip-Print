"""
Generate extra test samples:
1. sample_legacy.xls (BIFF8 .xls legacy workbook)
2. sample_multipage_invoice.xlsx (multi-page invoice with 35 items spanning 2-3 pages)
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

samples_dir = os.path.dirname(os.path.abspath(__file__))

def create_multipage_invoice():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Detailed Commercial Bill"
    
    ws.views.sheetView[0].showGridLines = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT

    # Set column widths
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 34
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 16

    title_font = Font(name="Segoe UI", size=16, bold=True, color="1E3A8A")
    regular_font = Font(name="Segoe UI", size=9, color="1F2937")
    bold_font = Font(name="Segoe UI", size=9, bold=True, color="111827")
    header_font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
    
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    alt_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )

    ws.merge_cells('A2:E2')
    ws['A2'] = "ENTERPRISE ANNUAL SERVICES BILL - MULTI-PAGE"
    ws['A2'].font = title_font
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')

    ws['A4'] = "Invoice #:"
    ws['A4'].font = bold_font
    ws['B4'] = "INV-2026-MULTI-099"
    ws['B4'].font = regular_font

    ws['D4'] = "Billing Date:"
    ws['D4'].font = bold_font
    ws['E4'] = "18-Oct-2026"
    ws['E4'].font = regular_font
    ws['E4'].alignment = Alignment(horizontal='right')

    headers = ["#", "Deliverable Milestone", "Hours", "Hourly Rate", "Amount (USD)"]
    ws.row_dimensions[6].height = 22
    for idx, (c, h) in enumerate(zip(['A', 'B', 'C', 'D', 'E'], headers)):
        cell = ws[f"{c}6"]
        cell.value = h
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center' if c in ['A', 'C'] else ('right' if c in ['D', 'E'] else 'left'), vertical='center')

    # Add 35 line items to force pagination across multiple pages
    curr_row = 7
    for i in range(1, 36):
        ws.row_dimensions[curr_row].height = 18
        ws[f"A{curr_row}"] = i
        ws[f"A{curr_row}"].alignment = Alignment(horizontal='center')
        ws[f"B{curr_row}"] = f"System Architecture Milestone Phase {i}: Module Deployment & Load Testing"
        ws[f"C{curr_row}"] = 8 + (i % 5)
        ws[f"C{curr_row}"].alignment = Alignment(horizontal='center')
        ws[f"D{curr_row}"] = 125.00
        ws[f"D{curr_row}"].number_format = "$#,##0.00"
        ws[f"D{curr_row}"].alignment = Alignment(horizontal='right')
        ws[f"E{curr_row}"] = f"=C{curr_row}*D{curr_row}"
        ws[f"E{curr_row}"].number_format = "$#,##0.00"
        ws[f"E{curr_row}"].alignment = Alignment(horizontal='right')

        for col in ['A', 'B', 'C', 'D', 'E']:
            ws[f"{col}{curr_row}"].border = thin_border
            ws[f"{col}{curr_row}"].font = regular_font
            if (curr_row % 2) == 1:
                ws[f"{col}{curr_row}"].fill = alt_fill

        curr_row += 1

    # Totals
    curr_row += 1
    ws[f"D{curr_row}"] = "Total Billed Amount:"
    ws[f"D{curr_row}"].font = Font(name="Segoe UI", size=10, bold=True, color="1E3A8A")
    ws[f"D{curr_row}"].alignment = Alignment(horizontal='right')
    ws[f"E{curr_row}"] = f"=SUM(E7:E{curr_row-2})"
    ws[f"E{curr_row}"].font = Font(name="Segoe UI", size=10, bold=True, color="1E3A8A")
    ws[f"E{curr_row}"].number_format = "$#,##0.00"
    ws[f"E{curr_row}"].alignment = Alignment(horizontal='right')

    # Signatory line
    curr_row += 3
    ws[f"D{curr_row}"] = "Authorized Executive Approval:"
    ws[f"D{curr_row}"].font = bold_font
    ws[f"D{curr_row}"].alignment = Alignment(horizontal='right')
    ws.merge_cells(f"D{curr_row}:E{curr_row}")

    ws[f"D{curr_row+3}"] = "------------------------------------------"
    ws[f"D{curr_row+3}"].font = regular_font
    ws[f"D{curr_row+3}"].alignment = Alignment(horizontal='right')
    ws.merge_cells(f"D{curr_row+3}:E{curr_row+3}")

    ws[f"D{curr_row+4}"] = "Alex Morgan, Managing Director"
    ws[f"D{curr_row+4}"].font = bold_font
    ws[f"D{curr_row+4}"].alignment = Alignment(horizontal='right')
    ws.merge_cells(f"D{curr_row+4}:E{curr_row+4}")

    out_path = os.path.join(samples_dir, "sample_multipage_invoice.xlsx")
    wb.save(out_path)
    print("Saved multipage invoice:", out_path)
    return out_path

def create_legacy_xls():
    """
    Creates an Excel 97-2003 BIFF8 .xls workbook via COM or xlwt.
    Using Excel COM since Excel is installed!
    """
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    excel = None
    wb = None
    out_path = os.path.join(samples_dir, "sample_legacy.xls")
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        wb = excel.Workbooks.Add()
        ws = wb.ActiveSheet
        ws.Name = "Legacy Bill"
        
        ws.Cells(2, 2).Value = "LEGACY ARCHIVE INVOICE (.XLS)"
        ws.Cells(2, 2).Font.Bold = True
        ws.Cells(2, 2).Font.Size = 14
        
        ws.Cells(4, 2).Value = "Bill No: LEGACY-9041"
        ws.Cells(4, 4).Value = "Date: 15-Oct-2026"
        
        headers = ["#", "Item Description", "Qty", "Rate", "Total"]
        for c, h in enumerate(headers, start=2):
            cell = ws.Cells(6, c)
            cell.Value = h
            cell.Font.Bold = True
            
        items = [
            (1, "Vintage Server Rack Hardware Maintenance", 2, 450.00),
            (2, "On-Site Tape Backup Archiving Service", 5, 120.00),
            (3, "Legacy Database Migration Protocol", 1, 1500.00),
        ]
        
        for r_idx, item in enumerate(items, start=7):
            num, desc, qty, rate = item
            ws.Cells(r_idx, 2).Value = num
            ws.Cells(r_idx, 3).Value = desc
            ws.Cells(r_idx, 4).Value = qty
            ws.Cells(r_idx, 5).Value = rate
            ws.Cells(r_idx, 6).Formula = f"=D{r_idx}*E{r_idx}"
            
        ws.Cells(11, 5).Value = "Total Due:"
        ws.Cells(11, 5).Font.Bold = True
        ws.Cells(11, 6).Formula = "=SUM(F7:F9)"
        ws.Cells(11, 6).Font.Bold = True
        
        # 56 = xlExcel8 (.xls)
        if os.path.exists(out_path):
            os.remove(out_path)
        wb.SaveAs(os.path.abspath(out_path), 56)
        print("Saved legacy .xls workbook:", out_path)
        return out_path
    finally:
        if wb:
            try:
                wb.Close(SaveChanges=False)
            except Exception:
                pass
        if excel:
            try:
                excel.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()

if __name__ == "__main__":
    create_multipage_invoice()
    create_legacy_xls()
