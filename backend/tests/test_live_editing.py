"""
Test live in-preview cell editing and text annotation features.
"""
import os
import sys
import unittest
import pymupdf as fitz

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

from backend.database import init_db
from backend.services.excel_converter import convert_excel_to_pdf
from backend.services.pdf_composer import compose_document
import openpyxl

class TestLiveEditing(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        init_db()
        cls.samples_dir = os.path.join(ROOT_DIR, "backend", "samples")
        cls.test_out_dir = os.path.join(ROOT_DIR, "backend", "storage", "test_outputs")
        cls.sample_excel = os.path.join(cls.samples_dir, "sample_invoice.xlsx")
        cls.letterhead_pdf = os.path.join(cls.samples_dir, "sample_letterhead.pdf")

    def test_live_cell_edit_reflected_in_pdf(self):
        # Create a working copy
        work_xlsx = os.path.join(self.test_out_dir, "live_edit_work.xlsx")
        import shutil
        shutil.copy2(self.sample_excel, work_xlsx)

        # 1. Edit cells
        wb = openpyxl.load_workbook(work_xlsx)
        ws = wb.active
        ws["B5"] = "INV-LIVE-EDITED-8888"
        ws["B8"] = "Worldwide Robotics Corporation"
        wb.save(work_xlsx)
        wb.close()

        # 2. Convert to raw PDF
        raw_pdf = os.path.join(self.test_out_dir, "live_edit_raw.pdf")
        convert_excel_to_pdf(work_xlsx, raw_pdf)

        # 3. Compose with letterhead and text annotation
        composed_pdf = os.path.join(self.test_out_dir, "live_edit_composed.pdf")
        annotations = [
            {
                "id": "ann-1",
                "text": "PAID VIA WIRE TRANSFER #7721",
                "x": 380,
                "y": 620,
                "width": 180,
                "height": 30,
                "font_size": 11,
                "font_color": "#16A34A",
                "is_bold": True,
                "page": "1",
                "rotation": 0
            }
        ]

        res = compose_document(
            bill_pdf_path=raw_pdf,
            template_path=self.letterhead_pdf,
            output_pdf_path=composed_pdf,
            annotations=annotations
        )
        self.assertTrue(res["success"])

        # 4. Verify text content in the output PDF
        doc = fitz.open(composed_pdf)
        text = doc[0].get_text()
        
        self.assertIn("INV-LIVE-EDITED-8888", text)
        self.assertIn("Worldwide Robotics Corporation", text)
        self.assertIn("PAID VIA WIRE TRANSFER #7721", text)
        doc.close()

        print("Test Live Editing Passed: Cell updates and text annotations are baked into the PDF!")

if __name__ == "__main__":
    unittest.main()
